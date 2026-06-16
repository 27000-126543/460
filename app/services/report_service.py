from sqlalchemy.orm import Session
from sqlalchemy.sql import func
from app.models import Booking, BookingStatus, WorkOrder, WorkOrderStatus, Resource, ResourceType, Equipment
from datetime import date, timedelta, datetime
from typing import Optional, List, Tuple
from openpyxl import Workbook
from io import BytesIO
import calendar


class ReportService:
    @staticmethod
    def get_daily_report(db: Session, report_date: date) -> dict:
        start_dt = datetime.combine(report_date, datetime.min.time())
        end_dt = datetime.combine(report_date, datetime.max.time())

        resources = db.query(Resource).all()
        resource_by_type_floor = {}
        for r in resources:
            key = (r.floor, r.type.value)
            if key not in resource_by_type_floor:
                resource_by_type_floor[key] = {"count": 0, "hourly_rate": r.hourly_rate}
            resource_by_type_floor[key]["count"] += 1

        bookings = db.query(Booking).filter(
            Booking.start_time <= end_dt,
            Booking.end_time >= start_dt,
            Booking.status.in_([BookingStatus.COMPLETED, BookingStatus.APPROVED])
        ).all()

        booked_hours_by_key = {}
        revenue_by_key = {}
        for booking in bookings:
            r = booking.resource
            if not r:
                continue
            key = (r.floor, r.type.value)
            
            overlap_start = max(booking.start_time, start_dt)
            overlap_end = min(booking.end_time, end_dt)
            hours = max(0, (overlap_end - overlap_start).total_seconds() / 3600)
            
            booked_hours_by_key[key] = booked_hours_by_key.get(key, 0) + hours
            
            amount = booking.actual_amount if booking.actual_amount else booking.total_amount
            revenue_by_key[key] = revenue_by_key.get(key, 0) + amount

        floor_utilizations = []
        total_revenue = 0
        total_booked_hours = 0
        total_available_hours = 0

        for (floor, rtype), info in resource_by_type_floor.items():
            count = info["count"]
            total_hours = count * 24
            booked_hours = booked_hours_by_key.get((floor, rtype), 0)
            revenue = revenue_by_key.get((floor, rtype), 0)
            utilization = booked_hours / total_hours if total_hours > 0 else 0

            floor_utilizations.append({
                "floor": floor,
                "resource_type": rtype,
                "total_resources": count,
                "total_hours": total_hours,
                "booked_hours": round(booked_hours, 2),
                "utilization_rate": round(utilization, 4),
                "revenue": round(revenue, 2)
            })

            total_revenue += revenue
            total_booked_hours += booked_hours
            total_available_hours += total_hours

        overall_utilization = total_booked_hours / total_available_hours if total_available_hours > 0 else 0

        work_orders = db.query(WorkOrder).filter(
            WorkOrder.created_at >= start_dt,
            WorkOrder.created_at <= end_dt
        ).all()

        total_orders = len(work_orders)
        completed_orders = sum(1 for w in work_orders if w.status == WorkOrderStatus.COMPLETED)

        response_times = []
        completion_times = []
        for w in work_orders:
            if w.accepted_at and w.created_at:
                response_times.append((w.accepted_at - w.created_at).total_seconds() / 60)
            if w.completed_at and w.created_at:
                completion_times.append((w.completed_at - w.created_at).total_seconds() / 60)

        avg_response = sum(response_times) / len(response_times) if response_times else 0
        avg_completion = sum(completion_times) / len(completion_times) if completion_times else 0

        data_available = total_booked_hours > 0 or total_orders > 0

        return {
            "report_date": report_date.isoformat(),
            "data_available": data_available,
            "floor_utilizations": sorted(floor_utilizations, key=lambda x: (x["floor"], x["resource_type"])),
            "total_revenue": round(total_revenue, 2),
            "overall_utilization_rate": round(overall_utilization, 4),
            "work_order_stats": {
                "total_orders": total_orders,
                "completed_orders": completed_orders,
                "avg_response_minutes": round(avg_response, 2),
                "avg_completion_minutes": round(avg_completion, 2)
            }
        }

    @staticmethod
    def _write_utilization_sheet(ws, floor_utilizations: list):
        ws.append(["楼层", "资源类型", "总资源数", "总时长(小时)", "已预订时长(小时)", "利用率(%)", "收入(元)"])
        for fu in floor_utilizations:
            ws.append([
                fu["floor"],
                fu["resource_type"],
                fu["total_resources"],
                fu["total_hours"],
                fu["booked_hours"],
                round(fu["utilization_rate"] * 100, 2),
                fu["revenue"]
            ])

    @staticmethod
    def _write_work_order_sheet(ws, work_order_stats: dict, total_revenue: float, overall_utilization: float):
        ws.append(["统计项", "数值"])
        ws.append(["总工单数", work_order_stats["total_orders"]])
        ws.append(["已完成数", work_order_stats["completed_orders"]])
        ws.append(["平均响应(分钟)", work_order_stats["avg_response_minutes"]])
        ws.append(["平均完成(分钟)", work_order_stats["avg_completion_minutes"]])
        ws.append(["总收入(元)", total_revenue])
        ws.append(["整体利用率(%)", round(overall_utilization * 100, 2)])

    @staticmethod
    def _write_trend_sheet(ws, trend_data: dict, title: str):
        ws.append(["周期", f"{title}"])
        for dp in trend_data["data_points"]:
            ws.append([dp["period"], dp["value"]])

    @staticmethod
    def _write_floor_comparison_sheet(ws, floor_comparisons: list):
        ws.append(["楼层", "总资源数", "总时长(小时)", "已预订时长(小时)", "利用率(%)", "收入(元)", "工单数"])
        for fc in floor_comparisons:
            ws.append([
                fc["floor"],
                fc["total_resources"],
                fc["total_hours"],
                fc["booked_hours"],
                round(fc["utilization_rate"] * 100, 2),
                fc["revenue"],
                fc["work_order_count"]
            ])

    @staticmethod
    def _write_resource_type_comparison_sheet(ws, resource_type_comparisons: list):
        ws.append(["资源类型", "总资源数", "总时长(小时)", "已预订时长(小时)", "利用率(%)", "收入(元)", "工单数"])
        for rc in resource_type_comparisons:
            ws.append([
                rc["resource_type"],
                rc["total_resources"],
                rc["total_hours"],
                rc["booked_hours"],
                round(rc["utilization_rate"] * 100, 2),
                rc["revenue"],
                rc["work_order_count"]
            ])

    @staticmethod
    def generate_excel_report(db: Session, report_date: date, report_type: str = 'daily') -> bytes:
        if report_type == 'daily':
            report = ReportService.get_daily_report(db, report_date)
        elif report_type == 'weekly':
            year, week = report_date
            report = ReportService.get_weekly_report(db, year, week)
        elif report_type == 'monthly':
            year, month = report_date
            report = ReportService.get_monthly_report(db, year, month)
        else:
            raise ValueError(f"Unsupported report_type: {report_type}")

        wb = Workbook()

        ws1 = wb.active
        ws1.title = "资源利用率"
        ReportService._write_utilization_sheet(ws1, report["floor_utilizations"])

        ws2 = wb.create_sheet("工单统计")
        ReportService._write_work_order_sheet(ws2, report["work_order_stats"], report["total_revenue"], report["overall_utilization_rate"])

        if report_type in ('weekly', 'monthly'):
            if report_type == 'weekly':
                year, week = report_date
                start_date = date.fromisocalendar(year, week, 1)
                end_date = date.fromisocalendar(year, week, 7)
                period = 'day'
            else:
                year, month = report_date
                _, last_day = calendar.monthrange(year, month)
                start_date = date(year, month, 1)
                end_date = date(year, month, last_day)
                period = 'day'

            ws3 = wb.create_sheet("趋势汇总")
            utilization_trend = ReportService.get_utilization_trend(db, start_date, end_date, period)
            revenue_trend = ReportService.get_revenue_trend(db, start_date, end_date, period)
            work_order_trend = ReportService.get_work_order_trend(db, start_date, end_date, period)

            ws3.append(["周期", "利用率(%)", "收入(元)", "工单数"])
            for i, dp in enumerate(utilization_trend["data_points"]):
                ws3.append([
                    dp["period"],
                    round(dp["value"] * 100, 2),
                    revenue_trend["data_points"][i]["value"] if i < len(revenue_trend["data_points"]) else 0,
                    work_order_trend["data_points"][i]["value"] if i < len(work_order_trend["data_points"]) else 0
                ])

            comparison = ReportService.get_comparison(db, start_date, start_date, end_date)

            ws4 = wb.create_sheet("楼层对比")
            ReportService._write_floor_comparison_sheet(ws4, comparison["floor_comparisons"])

            ws5 = wb.create_sheet("资源类型对比")
            ReportService._write_resource_type_comparison_sheet(ws5, comparison["resource_type_comparisons"])

        output = BytesIO()
        wb.save(output)
        return output.getvalue()

    @staticmethod
    def generate_weekly_excel_report(db: Session, year: int, week: int) -> bytes:
        return ReportService.generate_excel_report(db, (year, week), report_type='weekly')

    @staticmethod
    def generate_monthly_excel_report(db: Session, year: int, month: int) -> bytes:
        return ReportService.generate_excel_report(db, (year, month), report_type='monthly')

    @staticmethod
    def get_report_by_date_range(db: Session, start_date: date, end_date: date) -> List[dict]:
        reports = []
        current = start_date
        while current <= end_date:
            reports.append(ReportService.get_daily_report(db, current))
            current += timedelta(days=1)
        return reports

    @staticmethod
    def generate_booking_report(db, start_date=None, end_date=None):
        return {}

    @staticmethod
    def generate_payment_report(db, start_date=None, end_date=None):
        return {}

    @staticmethod
    def generate_equipment_report(db, start_date=None, end_date=None):
        return {}

    @staticmethod
    def _count_work_orders_by_dimension(db: Session, start_dt: datetime, end_dt: datetime, dimension: str) -> dict:
        if dimension == "floor":
            rows = (
                db.query(WorkOrder.floor, func.count(WorkOrder.id))
                .filter(WorkOrder.created_at >= start_dt, WorkOrder.created_at <= end_dt)
                .group_by(WorkOrder.floor)
                .all()
            )
            return {str(r[0]): r[1] for r in rows}
        elif dimension == "resource_type":
            from app.models import Equipment
            rows = (
                db.query(Equipment.type, func.count(WorkOrder.id))
                .join(Equipment, WorkOrder.equipment_id == Equipment.id)
                .filter(WorkOrder.created_at >= start_dt, WorkOrder.created_at <= end_dt)
                .group_by(Equipment.type)
                .all()
            )
            result = {}
            for r in rows:
                etype = r[0].value if hasattr(r[0], "value") else str(r[0])
                resource_type_mapping = {
                    "projector": "meeting_room",
                    "air_conditioner": "meeting_room",
                    "printer": "desk",
                    "router": "desk",
                    "light": "desk",
                    "other": "desk",
                }
                rtype = resource_type_mapping.get(etype, "desk")
                result[rtype] = result.get(rtype, 0) + r[1]
            return result
        elif dimension == "equipment_type":
            from app.models import Equipment
            rows = (
                db.query(Equipment.type, func.count(WorkOrder.id))
                .join(Equipment, WorkOrder.equipment_id == Equipment.id)
                .filter(WorkOrder.created_at >= start_dt, WorkOrder.created_at <= end_dt)
                .group_by(Equipment.type)
                .all()
            )
            return {r[0].value if hasattr(r[0], "value") else str(r[0]): r[1] for r in rows}
        return {}

    @staticmethod
    def _get_aggregated_report(db: Session, start_date: date, end_date: date) -> dict:
        start_dt = datetime.combine(start_date, datetime.min.time())
        end_dt = datetime.combine(end_date, datetime.max.time())
        days_count = (end_date - start_date).days + 1

        resources = db.query(Resource).all()
        resource_by_type_floor = {}
        for r in resources:
            key = (r.floor, r.type.value)
            if key not in resource_by_type_floor:
                resource_by_type_floor[key] = {"count": 0, "hourly_rate": r.hourly_rate}
            resource_by_type_floor[key]["count"] += 1

        bookings = db.query(Booking).filter(
            Booking.start_time <= end_dt,
            Booking.end_time >= start_dt,
            Booking.status.in_([BookingStatus.COMPLETED, BookingStatus.APPROVED])
        ).all()

        booked_hours_by_key = {}
        revenue_by_key = {}
        for booking in bookings:
            r = booking.resource
            if not r:
                continue
            key = (r.floor, r.type.value)

            overlap_start = max(booking.start_time, start_dt)
            overlap_end = min(booking.end_time, end_dt)
            hours = max(0, (overlap_end - overlap_start).total_seconds() / 3600)

            booked_hours_by_key[key] = booked_hours_by_key.get(key, 0) + hours

            amount = booking.actual_amount if booking.actual_amount else booking.total_amount
            revenue_by_key[key] = revenue_by_key.get(key, 0) + amount

        floor_utilizations = []
        total_revenue = 0
        total_booked_hours = 0
        total_available_hours = 0

        for (floor, rtype), info in resource_by_type_floor.items():
            count = info["count"]
            total_hours = count * 24 * days_count
            booked_hours = booked_hours_by_key.get((floor, rtype), 0)
            revenue = revenue_by_key.get((floor, rtype), 0)
            utilization = booked_hours / total_hours if total_hours > 0 else 0

            floor_utilizations.append({
                "floor": floor,
                "resource_type": rtype,
                "total_resources": count,
                "total_hours": total_hours,
                "booked_hours": round(booked_hours, 2),
                "utilization_rate": round(utilization, 4),
                "revenue": round(revenue, 2)
            })

            total_revenue += revenue
            total_booked_hours += booked_hours
            total_available_hours += total_hours

        overall_utilization = total_booked_hours / total_available_hours if total_available_hours > 0 else 0

        work_orders = db.query(WorkOrder).filter(
            WorkOrder.created_at >= start_dt,
            WorkOrder.created_at <= end_dt
        ).all()

        total_orders = len(work_orders)
        completed_orders = sum(1 for w in work_orders if w.status == WorkOrderStatus.COMPLETED)

        response_times = []
        completion_times = []
        for w in work_orders:
            if w.accepted_at and w.created_at:
                response_times.append((w.accepted_at - w.created_at).total_seconds() / 60)
            if w.completed_at and w.created_at:
                completion_times.append((w.completed_at - w.created_at).total_seconds() / 60)

        avg_response = sum(response_times) / len(response_times) if response_times else 0
        avg_completion = sum(completion_times) / len(completion_times) if completion_times else 0

        data_available = total_booked_hours > 0 or total_orders > 0

        return {
            "start_date": start_date.isoformat(),
            "end_date": end_date.isoformat(),
            "data_available": data_available,
            "floor_utilizations": sorted(floor_utilizations, key=lambda x: (x["floor"], x["resource_type"])),
            "total_revenue": round(total_revenue, 2),
            "overall_utilization_rate": round(overall_utilization, 4),
            "work_order_stats": {
                "total_orders": total_orders,
                "completed_orders": completed_orders,
                "avg_response_minutes": round(avg_response, 2),
                "avg_completion_minutes": round(avg_completion, 2)
            }
        }

    @staticmethod
    def get_weekly_report(db: Session, year: int, week: int) -> dict:
        try:
            start_date = date.fromisocalendar(year, week, 1)
            end_date = date.fromisocalendar(year, week, 7)
        except ValueError:
            return {
                "year": year,
                "week": week,
                "data_available": False,
                "floor_utilizations": [],
                "total_revenue": 0.0,
                "overall_utilization_rate": 0.0,
                "work_order_stats": {
                    "total_orders": 0,
                    "completed_orders": 0,
                    "avg_response_minutes": 0.0,
                    "avg_completion_minutes": 0.0
                }
            }

        report = ReportService._get_aggregated_report(db, start_date, end_date)
        return {
            "year": year,
            "week": week,
            "data_available": report["data_available"],
            "floor_utilizations": report["floor_utilizations"],
            "total_revenue": report["total_revenue"],
            "overall_utilization_rate": report["overall_utilization_rate"],
            "work_order_stats": report["work_order_stats"]
        }

    @staticmethod
    def get_monthly_report(db: Session, year: int, month: int) -> dict:
        try:
            _, last_day = calendar.monthrange(year, month)
            start_date = date(year, month, 1)
            end_date = date(year, month, last_day)
        except ValueError:
            return {
                "year": year,
                "month": month,
                "data_available": False,
                "floor_utilizations": [],
                "total_revenue": 0.0,
                "overall_utilization_rate": 0.0,
                "work_order_stats": {
                    "total_orders": 0,
                    "completed_orders": 0,
                    "avg_response_minutes": 0.0,
                    "avg_completion_minutes": 0.0
                }
            }

        report = ReportService._get_aggregated_report(db, start_date, end_date)
        return {
            "year": year,
            "month": month,
            "data_available": report["data_available"],
            "floor_utilizations": report["floor_utilizations"],
            "total_revenue": report["total_revenue"],
            "overall_utilization_rate": report["overall_utilization_rate"],
            "work_order_stats": report["work_order_stats"]
        }

    @staticmethod
    def _get_period_ranges(start_date: date, end_date: date, period: str) -> List[Tuple[date, date, str]]:
        ranges = []
        current = start_date

        if period == 'day':
            while current <= end_date:
                ranges.append((current, current, current.isoformat()))
                current += timedelta(days=1)
        elif period == 'week':
            while current <= end_date:
                week_start = current - timedelta(days=current.weekday())
                week_end = week_start + timedelta(days=6)
                if week_end > end_date:
                    week_end = end_date
                label = f"{week_start.isoformat()} ~ {week_end.isoformat()}"
                ranges.append((week_start, week_end, label))
                current = week_end + timedelta(days=1)
        elif period == 'month':
            while current <= end_date:
                month_start = date(current.year, current.month, 1)
                _, last_day = calendar.monthrange(current.year, current.month)
                month_end = date(current.year, current.month, last_day)
                if month_end > end_date:
                    month_end = end_date
                label = f"{current.year}-{current.month:02d}"
                ranges.append((month_start, month_end, label))
                current = month_end + timedelta(days=1)
        else:
            raise ValueError(f"Unsupported period: {period}")

        return ranges

    @staticmethod
    def get_utilization_trend(db: Session, start_date: date, end_date: date, period: str = 'day') -> dict:
        ranges = ReportService._get_period_ranges(start_date, end_date, period)
        data_points = []

        for period_start, period_end, label in ranges:
            report = ReportService._get_aggregated_report(db, period_start, period_end)
            data_points.append({
                "period": label,
                "value": report["overall_utilization_rate"]
            })

        return {
            "start_date": start_date.isoformat(),
            "end_date": end_date.isoformat(),
            "period": period,
            "data_points": data_points
        }

    @staticmethod
    def get_revenue_trend(db: Session, start_date: date, end_date: date, period: str = 'day') -> dict:
        ranges = ReportService._get_period_ranges(start_date, end_date, period)
        data_points = []

        for period_start, period_end, label in ranges:
            report = ReportService._get_aggregated_report(db, period_start, period_end)
            data_points.append({
                "period": label,
                "value": report["total_revenue"]
            })

        return {
            "start_date": start_date.isoformat(),
            "end_date": end_date.isoformat(),
            "period": period,
            "data_points": data_points
        }

    @staticmethod
    def get_work_order_trend(db: Session, start_date: date, end_date: date, period: str = 'day') -> dict:
        ranges = ReportService._get_period_ranges(start_date, end_date, period)
        data_points = []

        for period_start, period_end, label in ranges:
            report = ReportService._get_aggregated_report(db, period_start, period_end)
            data_points.append({
                "period": label,
                "value": float(report["work_order_stats"]["total_orders"])
            })

        return {
            "start_date": start_date.isoformat(),
            "end_date": end_date.isoformat(),
            "period": period,
            "data_points": data_points
        }

    @staticmethod
    def get_floor_comparison(db: Session, report_date: date, start_date: Optional[date] = None, end_date: Optional[date] = None) -> dict:
        if start_date and end_date:
            actual_start = start_date
            actual_end = end_date
        else:
            actual_start = report_date
            actual_end = report_date

        report = ReportService._get_aggregated_report(db, actual_start, actual_end)

        start_dt = datetime.combine(actual_start, datetime.min.time())
        end_dt = datetime.combine(actual_end, datetime.max.time())
        floor_wo_counts = ReportService._count_work_orders_by_dimension(db, start_dt, end_dt, "floor")

        floor_data = {}
        for fu in report["floor_utilizations"]:
            floor = fu["floor"]
            if floor not in floor_data:
                floor_data[floor] = {
                    "floor": floor,
                    "total_resources": 0,
                    "total_hours": 0.0,
                    "booked_hours": 0.0,
                    "revenue": 0.0
                }
            floor_data[floor]["total_resources"] += fu["total_resources"]
            floor_data[floor]["total_hours"] += fu["total_hours"]
            floor_data[floor]["booked_hours"] += fu["booked_hours"]
            floor_data[floor]["revenue"] += fu["revenue"]

        floor_comparisons = []
        for floor, data in floor_data.items():
            utilization = data["booked_hours"] / data["total_hours"] if data["total_hours"] > 0 else 0
            floor_comparisons.append({
                "floor": floor,
                "total_resources": data["total_resources"],
                "total_hours": data["total_hours"],
                "booked_hours": round(data["booked_hours"], 2),
                "utilization_rate": round(utilization, 4),
                "revenue": round(data["revenue"], 2),
                "work_order_count": floor_wo_counts.get(str(floor), 0)
            })

        return {
            "report_date": report_date.isoformat(),
            "floor_comparisons": sorted(floor_comparisons, key=lambda x: x["floor"]),
            "resource_type_comparisons": []
        }

    @staticmethod
    def get_resource_type_comparison(db: Session, report_date: date, start_date: Optional[date] = None, end_date: Optional[date] = None) -> dict:
        if start_date and end_date:
            actual_start = start_date
            actual_end = end_date
        else:
            actual_start = report_date
            actual_end = report_date

        report = ReportService._get_aggregated_report(db, actual_start, actual_end)

        start_dt = datetime.combine(actual_start, datetime.min.time())
        end_dt = datetime.combine(actual_end, datetime.max.time())
        rtype_wo_counts = ReportService._count_work_orders_by_dimension(db, start_dt, end_dt, "resource_type")

        type_data = {}
        for fu in report["floor_utilizations"]:
            rtype = fu["resource_type"]
            if rtype not in type_data:
                type_data[rtype] = {
                    "resource_type": rtype,
                    "total_resources": 0,
                    "total_hours": 0.0,
                    "booked_hours": 0.0,
                    "revenue": 0.0
                }
            type_data[rtype]["total_resources"] += fu["total_resources"]
            type_data[rtype]["total_hours"] += fu["total_hours"]
            type_data[rtype]["booked_hours"] += fu["booked_hours"]
            type_data[rtype]["revenue"] += fu["revenue"]

        resource_type_comparisons = []
        for rtype, data in type_data.items():
            utilization = data["booked_hours"] / data["total_hours"] if data["total_hours"] > 0 else 0
            resource_type_comparisons.append({
                "resource_type": rtype,
                "total_resources": data["total_resources"],
                "total_hours": data["total_hours"],
                "booked_hours": round(data["booked_hours"], 2),
                "utilization_rate": round(utilization, 4),
                "revenue": round(data["revenue"], 2),
                "work_order_count": rtype_wo_counts.get(rtype, 0)
            })

        return {
            "report_date": report_date.isoformat(),
            "floor_comparisons": [],
            "resource_type_comparisons": sorted(resource_type_comparisons, key=lambda x: x["resource_type"])
        }

    @staticmethod
    def get_comparison(db: Session, report_date: date, start_date: Optional[date] = None, end_date: Optional[date] = None) -> dict:
        if start_date and end_date:
            actual_start = start_date
            actual_end = end_date
        else:
            actual_start = report_date
            actual_end = report_date

        report = ReportService._get_aggregated_report(db, actual_start, actual_end)

        start_dt = datetime.combine(actual_start, datetime.min.time())
        end_dt = datetime.combine(actual_end, datetime.max.time())
        floor_wo_counts = ReportService._count_work_orders_by_dimension(db, start_dt, end_dt, "floor")
        rtype_wo_counts = ReportService._count_work_orders_by_dimension(db, start_dt, end_dt, "resource_type")

        floor_data = {}
        type_data = {}
        for fu in report["floor_utilizations"]:
            floor = fu["floor"]
            rtype = fu["resource_type"]

            if floor not in floor_data:
                floor_data[floor] = {
                    "floor": floor,
                    "total_resources": 0,
                    "total_hours": 0.0,
                    "booked_hours": 0.0,
                    "revenue": 0.0
                }
            floor_data[floor]["total_resources"] += fu["total_resources"]
            floor_data[floor]["total_hours"] += fu["total_hours"]
            floor_data[floor]["booked_hours"] += fu["booked_hours"]
            floor_data[floor]["revenue"] += fu["revenue"]

            if rtype not in type_data:
                type_data[rtype] = {
                    "resource_type": rtype,
                    "total_resources": 0,
                    "total_hours": 0.0,
                    "booked_hours": 0.0,
                    "revenue": 0.0
                }
            type_data[rtype]["total_resources"] += fu["total_resources"]
            type_data[rtype]["total_hours"] += fu["total_hours"]
            type_data[rtype]["booked_hours"] += fu["booked_hours"]
            type_data[rtype]["revenue"] += fu["revenue"]

        floor_comparisons = []
        for floor, data in floor_data.items():
            utilization = data["booked_hours"] / data["total_hours"] if data["total_hours"] > 0 else 0
            floor_comparisons.append({
                "floor": floor,
                "total_resources": data["total_resources"],
                "total_hours": data["total_hours"],
                "booked_hours": round(data["booked_hours"], 2),
                "utilization_rate": round(utilization, 4),
                "revenue": round(data["revenue"], 2),
                "work_order_count": floor_wo_counts.get(str(floor), 0)
            })

        resource_type_comparisons = []
        for rtype, data in type_data.items():
            utilization = data["booked_hours"] / data["total_hours"] if data["total_hours"] > 0 else 0
            resource_type_comparisons.append({
                "resource_type": rtype,
                "total_resources": data["total_resources"],
                "total_hours": data["total_hours"],
                "booked_hours": round(data["booked_hours"], 2),
                "utilization_rate": round(utilization, 4),
                "revenue": round(data["revenue"], 2),
                "work_order_count": rtype_wo_counts.get(rtype, 0)
            })

        return {
            "report_date": report_date.isoformat(),
            "floor_comparisons": sorted(floor_comparisons, key=lambda x: x["floor"]),
            "resource_type_comparisons": sorted(resource_type_comparisons, key=lambda x: x["resource_type"])
        }
